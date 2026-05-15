import csv
from pathlib import Path

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = PROJECT_ROOT / "01_sources" / "heim_primary" / "Heim_1989_Massenformel_0.4.xlsm"
OUT_DIR = PROJECT_ROOT / "07_outputs" / "xlsm_static_export"


def clean_name(name: str) -> str:
    return "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in name)


def export_workbook(data_only: bool, suffix: str) -> None:
    workbook = openpyxl.load_workbook(
        WORKBOOK,
        read_only=True,
        data_only=data_only,
        keep_vba=False,
    )

    for sheet in workbook.worksheets:
        path = OUT_DIR / f"{clean_name(sheet.title)}_{suffix}.csv"
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            for row in sheet.iter_rows(values_only=data_only):
                if data_only:
                    writer.writerow(["" if value is None else value for value in row])
                else:
                    writer.writerow(
                        ["" if cell is None or cell.value is None else cell.value for cell in row]
                    )


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    export_workbook(data_only=True, suffix="cached_values")
    export_workbook(data_only=False, suffix="formulas")

    workbook = openpyxl.load_workbook(
        WORKBOOK,
        read_only=True,
        data_only=False,
        keep_vba=False,
    )
    summary_path = OUT_DIR / "workbook_summary.csv"
    with summary_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["sheet", "max_row", "max_column"])
        for sheet in workbook.worksheets:
            writer.writerow([sheet.title, sheet.max_row, sheet.max_column])


if __name__ == "__main__":
    main()
